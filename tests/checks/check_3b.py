import os, sys, tempfile, time, io, dataclasses
sys.path.insert(0, os.getcwd())
tmp = tempfile.mkdtemp()
import chromadb
from fastapi import FastAPI
from fastapi.testclient import TestClient
from app import retriever, global_settings, config
from app.admin import api as A, users, kbs_api, documents_api
from app.admin.auth import SessionSigner
from app.channels.manager import manager
from app.channels.store import ChannelStore
from app.db import Database
from app.knowledge import kbs, documents, metadata
from app.secretstore import FernetSecretStore, generate_key

db = Database(os.path.join(tmp, "t.db")); db.migrate()
client = chromadb.PersistentClient(path=os.path.join(tmp, "chroma"))
retriever.get_client = lambda: client
kbs.get_db = documents.get_db = global_settings.get_db = lambda: db
cfg = dataclasses.replace(config.settings, docs_dir=os.path.join(tmp, "files"), max_upload_mb=1)
documents.settings = documents_api.settings = cfg
documents.RETRY_WAIT = (0.01, 0.05)
for _m in list(sys.modules.values()):                     # every module that reads the database reads this one, whatever it imported
    if getattr(_m, "__name__", "").startswith("app") and hasattr(_m, "get_db"):
        _m.get_db = lambda: db
ok = lambda cond, msg: print(("PASS " if cond else "FAIL ") + msg)

# legacy documents already in the index, with the original still in data/docs
default = kbs.ensure_default(db)
col = retriever.get_collection()
col.upsert(ids=["HR-Policy.pdf::0"], documents=["Earned leave is 18 days."], metadatas=[{"source": "HR-Policy.pdf", "doc_id": "HR-Policy.pdf", "page": 1}])
kbs.backfill_metadata(db, default)
ok(documents.import_existing(db) == 1 and documents.import_existing(db) == 0, "existing indexed document recorded once")
legacy = documents.list_for_kb(db, default["id"])[0]
ok(legacy["name"] == "HR-Policy.pdf" and legacy["status"] == "Indexed" and documents.file_path(legacy).exists(), "legacy record keeps name, status and a copy of the original")

store = ChannelStore(db, FernetSecretStore(generate_key()), lambda t: manager.adapter_class(t).SPEC)
ctx = A.AdminContext(db=db, store=store, manager=manager, signer=SessionSigner(generate_key()), domain="sysfore.com", allow_remote=True)
users.create_admin(db, "boss@sysfore.com", "correct horse battery", domain="sysfore.com")
app = FastAPI()
for r in (A.router, kbs_api.router, documents_api.router):
    app.include_router(r)
app.dependency_overrides[A.get_ctx] = lambda: ctx
c = TestClient(app, base_url="http://localhost")
r = c.post("/admin/api/auth/login", json={"email": "boss@sysfore.com", "password": "correct horse battery"})
h = {"X-CSRF-Token": r.json()["csrf"]}
it = c.post("/admin/api/kbs", headers=h, json={"name": "IT help", "department_id": default["department_id"]}).json()


def wait(doc_id, want=("Indexed", "Failed"), t=60):
    end = time.time() + t
    while time.time() < end:
        d = documents.get(db, doc_id)
        if d["status"] in want:
            return d
        time.sleep(0.2)
    return documents.get(db, doc_id)


def up(kb, name, data):
    return c.post(f"/admin/api/kbs/{kb}/documents", headers=h, files={"file": (name, data)})


# build one file of each type
from docx import Document as W
w = W(); w.add_paragraph("Laptop refresh happens every three years for all employees."); b = io.BytesIO(); w.save(b); docx = b.getvalue()
from pptx import Presentation
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "VPN"; s.placeholders[1].text = "Reset your VPN token via the service desk."
b = io.BytesIO(); p.save(b); pptx = b.getvalue()
pdf = open("data/docs/HR-Policy.pdf", "rb").read()

results = {}
for name, data in (("HR-Policy.pdf", pdf), ("laptops.docx", docx), ("vpn.pptx", pptx), ("notes.md", b"# Notes\nParking is free for staff.")):
    r = up(it["id"], name, data)
    ok(r.status_code == 201 and r.json()["status"] in ("Pending", "Indexing", "Indexed"), f"{name} accepted")
    results[name] = r.json()["id"]
for name, did in results.items():
    d = wait(did)
    ok(d["status"] == "Indexed" and d["chunks"] > 0, f"{name}: Pending -> Indexed ({d['chunks']} chunks) {d['error']}")
hits = retriever.retrieve_scored("how long is the laptop refresh cycle", kb_ids=[it["id"]])
ok(hits and hits[0]["source"] == "laptops.docx", "docx content is searchable")
hits = retriever.retrieve_scored("reset VPN token", kb_ids=[it["id"]])
ok(hits and hits[0]["source"] == "vpn.pptx" and hits[0]["page"] == 1, "pptx content is searchable, slide number as page")
hits = retriever.retrieve_scored("earned leave days", kb_ids=[it["id"]])
ok(hits and hits[0]["source"] == "HR-Policy.pdf" and hits[0]["page"], "pdf cites file and page")
meta = [m for _, ms in kbs.store_for(kbs.get_kb(db, it["id"])).iter_metadata() for m in ms]
ok(all(not metadata.missing(m) and m["kb_id"] == it["id"] for m in meta), "every chunk has the full metadata")

# validation
ok(up(it["id"], "notes.md", b"again").status_code == 409, "duplicate name refused (use Replace)")
ok(up(it["id"], "big.txt", b"x" * (1024 * 1024 + 10)).status_code in (413, 422), "over the size limit refused")
ok(up(it["id"], "old.doc", b"PK\x03\x04x").status_code == 422, "old .doc refused with a hint")
ok(up(it["id"], "fake.pdf", b"not a pdf").status_code == 422, "file that is not a PDF refused")
ok(up(it["id"], "empty.txt", b"").status_code == 422, "empty file refused")
ev = up(it["id"], "../../evil.txt", b"hello").json()
wait(ev["id"])
ok(ev["name"] == "evil.txt", "path tricks in the file name are neutralised")

# forced embedding failure -> Failed with error; retry works for that document alone
real_add = kbs.ChromaStore.add
state = {"fail": True, "calls": 0}


def flaky(self, ids, texts, metas):
    state["calls"] += 1
    if state["fail"]:
        raise RuntimeError("embedding service unavailable")
    return real_add(self, ids, texts, metas)


kbs.ChromaStore.add = flaky
r = up(it["id"], "fail-me.txt", b"A file that cannot be embedded right now.")
d = wait(r.json()["id"])
ok(d["status"] == "Failed" and "unavailable" in d["error"] and state["calls"] == documents.RETRY_ATTEMPTS, f"forced failure -> Failed after {state['calls']} attempts, error shown")
others = [documents.get(db, i)["status"] for i in results.values()]
ok(set(others) == {"Indexed"}, "other documents untouched by the failure")
state["fail"] = False
c.post(f"/admin/api/documents/{d['id']}/reindex", headers=h)
d = wait(d["id"])
ok(d["status"] == "Indexed" and not d["error"], "single-document retry works")

# replace, move, delete
rid = results["notes.md"]
r = c.post(f"/admin/api/documents/{rid}/replace", headers=h, files={"file": ("notes.md", b"# Notes\nCanteen opens at 8am.")})
d = wait(rid)
ok(r.status_code == 200 and d["status"] == "Indexed" and retriever.retrieve_scored("canteen opens", kb_ids=[it["id"]])[0]["source"] == "notes.md", "replace re-indexes the new content")
ok(not any("Parking" in x["text"] for x in retriever.retrieve_scored("parking free", kb_ids=[it["id"]]) if x["source"] == "notes.md"), "old content of the replaced file is gone")
r = c.post(f"/admin/api/documents/{rid}/move", headers=h, json={"kb_id": default["id"]})
d = wait(rid)
ok(d["kb_id"] == default["id"] and d["status"] == "Indexed", "move re-indexes in the target KB")
ok(not [x for x in retriever.retrieve_scored("canteen opens", kb_ids=[it["id"]]) if x["source"] == "notes.md"], "nothing of it left in the old KB")
ok(retriever.retrieve_scored("canteen opens")[0]["source"] == "notes.md", "found in the target KB")
c.delete(f"/admin/api/documents/{rid}", headers=h)
ok(documents.get(db, rid) is None and not documents.doc_dir(default["id"], rid).exists() and not [x for x in retriever.retrieve_scored("canteen opens") if x["source"] == "notes.md"], "delete removes record, file and vectors")

# re-index the whole KB, and KB delete leaves nothing behind
q = c.post(f"/admin/api/kbs/{it['id']}/reindex", headers=h).json()["queued"]
ok(q >= 4, f"re-index all queued {q} documents")
for d in documents.list_for_kb(db, it["id"]):
    wait(d["id"])
r = c.delete(f"/admin/api/kbs/{it['id']}", headers=h).json()
ok(r["documents_removed"] >= 4 and not (os.path.exists(os.path.join(cfg.docs_dir, str(it["id"])))), "deleting the KB removes its documents and files")
ok(documents.list_for_kb(db, it["id"]) == [] and it["id"] not in [k["id"] for k in kbs.list_kbs(db)], "no orphan document rows")
documents.shutdown()
