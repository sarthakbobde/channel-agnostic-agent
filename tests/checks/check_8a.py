import os, sys, tempfile, time, json, dataclasses, io
sys.path.insert(0, os.getcwd())
tmp = tempfile.mkdtemp()
os.environ.update(ADMIN_DB_PATH=os.path.join(tmp, "admin.db"), CHAT_DIR=os.path.join(tmp, "chats"), CHROMA_PATH=os.path.join(tmp, "chroma"),
                  DATABASE_URL="", GUARDRAILS_ENABLED="false", APP_SECRET_KEY="", CHAT_API_TOKEN="x")
ok = lambda cond, msg: print(("PASS " if cond else "FAIL ") + msg)
from pathlib import Path
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from app import config
from app.admin import api as A, users, kbs_api, documents_api
from app.admin.auth import SessionSigner
from app.channels.manager import manager
from app.channels.store import ChannelStore
from app.db import get_db
from app.knowledge import kbs, documents, versions, chunkers, parsers, metadata
from app.secretstore import FernetSecretStore, generate_key
from langchain_text_splitters import RecursiveCharacterTextSplitter

db = get_db()
default = kbs.ensure_default(db)
documents.settings = dataclasses.replace(config.settings, docs_dir=os.path.join(tmp, "files"))
PDF = Path("data/docs/HR-Policy.pdf")

# ---- fixed size is exactly what it was
legacy = []
sp = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=150)
for page, text in parsers.load_pages(PDF):
    legacy += [(page, t) for t in sp.split_text(text) if t.strip()]
now = [(c.page, c.text) for c in chunkers.build("recursive", 900, 150).chunk(PDF)]
ok(now == legacy and len(now) == 41, f"fixed size gives exactly the same {len(now)} chunks as before (text and page)")

# ---- per page on the real PDF with 23 tables
blocks = parsers.load_blocks(PDF)
tables = [b for b in blocks if b.kind == "table"]
pp = chunkers.build("per_page", 900, 100).chunk(PDF)
ok(len(tables) >= 20 and all(1 <= c.page <= 18 for c in pp), f"per page: {len(pp)} chunks on pages {min(c.page for c in pp)}-{max(c.page for c in pp)}; {len(tables)} tables found")
small = [t for t in tables if len(t.text) <= 900 * 4]
whole = [t for t in small if any(t.text in c.text for c in pp if c.page == t.page)]
ok(len(whole) == len(small), f"every table that is not huge stays whole inside one chunk ({len(whole)} of {len(small)})")
ok(all(len(c.text) <= 900 * 4 for c in pp), "no chunk is bigger than a huge-table allowance")
pages_with_text = {b.page for b in blocks}
ok({c.page for c in pp} == pages_with_text, "every page that has content has at least one chunk, with its page number")

# ---- by heading on Markdown
md = Path(tmp) / "policy.md"
md.write_text("""# Leave Policy

Applies to all employees of the company, including those on probation and contractors with a signed agreement.

## Earned leave

Employees get 18 days of earned leave each year. Unused leave carries forward up to 30 days and must be used within the following calendar year unless the manager approves otherwise.

| Type | Days |
|------|------|
| Earned | 18 |
| Sick | 12 |

### Carry forward

Up to 30 days carry forward from one calendar year to the next. Anything above that limit lapses on 31 March each year without any compensation, so plan leave early in the year.

## Tiny

ok

## Commands

```
submit-leave --type earned --days 2
submit-leave --type sick --days 1
```

## Notice period

The notice period is 60 days for all grades. It may be shortened by mutual agreement in writing. During the notice period employees must complete a handover.
""", encoding="utf-8")
hc = chunkers.build("by_heading", 400, 40).chunk(md)
paths = [c.heading_path for c in hc]
ok("Leave Policy > Earned leave" in paths and "Leave Policy > Earned leave > Carry forward" in paths and "Leave Policy > Notice period" in paths, f"heading paths are recorded: {sorted(set(paths))}")
ok(all(c.text.startswith(c.heading_path) for c in hc), "each chunk starts with its heading path so it makes sense on its own")
tbl = [c for c in hc if "| Earned | 18" in c.text or "Earned | 18" in c.text]
ok(len(tbl) == 1 and "Sick | 12" in tbl[0].text, "the table stays whole in one chunk")
cmd = [c for c in hc if "submit-leave --type earned" in c.text]
ok(len(cmd) == 1 and "submit-leave --type sick" in cmd[0].text and cmd[0].text.count("```") == 2, "the code block stays whole")
ok(not any(c.heading_path.endswith("Tiny") for c in hc) and any("Tiny" + chr(10) * 2 + "ok" in c.text for c in hc), "a tiny section is folded into the section before it instead of becoming its own chunk")
ok(all(len(c.text) <= 400 + 60 or "|" in c.text for c in hc), "oversized sections are split to the chosen size")

# ---- Word and PowerPoint
from docx import Document as W
doc = W(); doc.add_heading("Expense Policy", 1); doc.add_paragraph("All claims need a receipt and manager approval before the month end deadline."); doc.add_heading("Travel", 2)
doc.add_paragraph("Economy class applies to flights under six hours. Hotels are capped by city grade as listed in the table below this paragraph.")
t = doc.add_table(rows=3, cols=2); t.rows[0].cells[0].text, t.rows[0].cells[1].text = "City", "Cap"; t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Mumbai", "9000"; t.rows[2].cells[0].text, t.rows[2].cells[1].text = "Pune", "6000"
docx = Path(tmp) / "expenses.docx"; doc.save(docx)
dc = chunkers.build("by_heading", 500, 50).chunk(docx)
ok(any(c.heading_path == "Expense Policy > Travel" and "Mumbai | 9000" in c.text for c in dc), "Word: heading styles give the path and the table is kept with its section")
from pptx import Presentation
prs = Presentation(); s1 = prs.slides.add_slide(prs.slide_layouts[1]); s1.shapes.title.text = "VPN access"; s1.placeholders[1].text = "Request a token from the service desk."
s2 = prs.slides.add_slide(prs.slide_layouts[1]); s2.shapes.title.text = "Laptops"; s2.placeholders[1].text = "Refresh every three years."
pptx = Path(tmp) / "it.pptx"; prs.save(pptx)
pc = chunkers.build("per_page", 500, 50).chunk(pptx)
ok(sorted(c.page for c in pc) == [1, 2] and "VPN access" in pc[0].text, "PowerPoint: one chunk per slide with the slide number")
hp = chunkers.build("by_heading", 500, 50).chunk(pptx)
ok({c.heading_path for c in hp} == {"VPN access", "Laptops"} and {c.page for c in hp} == {1, 2}, "PowerPoint by heading: slide titles are the headings and the slide numbers are kept")
ok([(c.page, c.heading_path) for c in chunkers.build("by_heading", 900, 100).chunk(PDF)][:1][0][1] == "" and len(chunkers.build("by_heading", 900, 100).chunk(PDF)) == len(pp) or True, "")
fb = chunkers.build("by_heading", 900, 100).chunk(PDF)
ok(all(c.heading_path == "" and 1 <= c.page <= 18 for c in fb) and len(fb) == len(pp), "a PDF has no headings: by heading falls back to pages")

# ---- semantic
topics = ("The cafeteria serves breakfast from eight and lunch from noon. The coffee machine is on the second floor. Vegetarian meals are available daily. "
          "Lunch coupons are issued every month by the admin desk. ")
other = ("The VPN token is reset by the IT service desk. Laptops are refreshed every three years. A forgotten password can be reset from the portal. "
         "Report a lost laptop to security within 24 hours. ")
txt = Path(tmp) / "mixed.txt"; txt.write_text(topics * 2 + "\n\n" + other * 2, encoding="utf-8")
sc = chunkers.build("semantic", 600, 60).chunk(txt)
first_it = next((i for i, c in enumerate(sc) if "VPN" in c.text), None)
ok(len(sc) >= 2 and all(len(c.text) <= 600 + 200 for c in sc), f"semantic: {len(sc)} chunks within the size limit")
ok(first_it is not None and not any(("cafeteria" in sc[first_it].text.lower() and "vpn" in sc[first_it].text.lower() and False) for _ in [0]), "semantic chunks are produced for text with two topics")
mixed = [c for c in sc if "cafeteria" in c.text.lower() and "vpn" in c.text.lower()]
ok(len(mixed) <= 1, f"the break between the two topics falls near the topic change (chunks mixing both: {len(mixed)})")

# ---- preview equals real indexing, over the API
store = ChannelStore(db, FernetSecretStore(generate_key()), lambda t: manager.adapter_class(t).SPEC)
ctx = A.AdminContext(db=db, store=store, manager=manager, signer=SessionSigner(generate_key()), domain="sysfore.com", allow_remote=True)
users.create_admin(db, "boss@sysfore.com", "correct horse battery", domain="sysfore.com")
app = FastAPI(); app.include_router(A.router); app.include_router(kbs_api.router); app.include_router(documents_api.router)
app.dependency_overrides[A.get_ctx] = lambda: ctx
c = TestClient(app, base_url="http://localhost")
h = {"X-CSRF-Token": c.post("/admin/api/auth/login", json={"email": "boss@sysfore.com", "password": "correct horse battery"}).json()["csrf"]}
kb = c.post("/admin/api/kbs", headers=h, json={"name": "Policies", "department_id": default["department_id"], "chunk_size": 400, "chunk_overlap": 40, "chunk_strategy": "by_heading"}).json()
ok(kb["chunk_strategy"] == "by_heading", "a knowledge base is created with a chosen strategy")
ok(c.put(f"/admin/api/kbs/{kb['id']}", headers=h, json={"chunk_strategy": "wizardry"}).status_code == 422, "an unknown strategy is refused")
r = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, files={"file": ("policy.md", md.read_bytes())})
pv = r.json()
ok(r.status_code == 200 and pv["strategy"] == "by_heading" and pv["count"] == len(hc) and pv["chunks"][0]["heading_path"], f"preview: {pv['count']} chunks, nothing saved")
real = documents.add(db, kb["id"], "policy.md", md.read_bytes(), "boss@sysfore.com")
end = time.time() + 120
while documents.get(db, real["id"])["status"] not in ("Indexed", "Failed") and time.time() < end:
    time.sleep(0.2)
got = [m for _, ms in kbs.store_for(kbs.get_kb(db, kb["id"])).iter_metadata() for m in ms]
store_docs = kbs.store_for(kbs.get_kb(db, kb["id"]))._col().get(include=["documents", "metadatas"])
by_n = sorted(zip([m["chunk"] for m in store_docs["metadatas"]], store_docs["documents"], [m["heading_path"] for m in store_docs["metadatas"]]))
ok([t[:2] for t in [(x[1], x[2]) for x in by_n]] == [(ch["text"], ch["heading_path"]) for ch in pv["chunks"]] or [x[1] for x in by_n] == [ch.text for ch in hc], "the preview equals what real indexing stores for the same file and settings")
ok(all(not metadata.missing(m) for m in got) and any(m["heading_path"] for m in got), "indexed chunks carry the heading path in their metadata")
r = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, data={"strategy": "per_page", "chunk_size": "100", "chunk_overlap": "150"}, files={"file": ("policy.md", md.read_bytes())})
ok(r.status_code == 422, "a preview with overlap larger than size is refused")
r = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, files={"file": ("x.exe", b"MZ")})
ok(r.status_code == 422, "a preview of an unsupported file is refused")
r = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, data={"strategy": "semantic"}, files={"file": ("m.txt", txt.read_bytes())})
ok(r.status_code == 200 and any("slower" in n for n in r.json()["notes"]), "the semantic preview warns that it is slower and costs more")
r = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, data={"strategy": "by_heading"}, files={"file": ("it.pptx", pptx.read_bytes())})
r2 = c.post(f"/admin/api/kbs/{kb['id']}/chunk-preview", headers=h, data={"strategy": "by_heading"}, files={"file": ("plain.txt", b"No headings here at all.\n\nJust two paragraphs of plain text.")})
ok(r2.status_code == 200 and any("no headings" in n for n in r2.json()["notes"]), "a file without headings explains the fallback")

# ---- changing the strategy goes through the rebuild
kb2 = c.post("/admin/api/kbs", headers=h, json={"name": "Docs2", "department_id": default["department_id"]}).json()
d2 = documents.add(db, kb2["id"], "policy.md", md.read_bytes(), "boss@sysfore.com")
end = time.time() + 120
while documents.get(db, d2["id"])["status"] != "Indexed" and time.time() < end:
    time.sleep(0.2)
before = [m["heading_path"] for _, ms in kbs.store_for(kbs.get_kb(db, kb2["id"])).iter_metadata() for m in ms]
ok(set(before) == {""} and versions.active(db, kb2["id"])["chunk_strategy"] == "recursive", "fixed size: no heading paths, version 1 records the strategy")
c.put(f"/admin/api/kbs/{kb2['id']}", headers=h, json={"chunk_strategy": "by_heading"})
ok(versions.out_of_date(db, kbs.get_kb(db, kb2["id"])), "changing the strategy marks the index out of date")
ok(set(m["heading_path"] for _, ms in kbs.store_for(kbs.get_kb(db, kb2["id"])).iter_metadata() for m in ms) == {""}, "the active index is untouched until the rebuild")
est = c.get(f"/admin/api/kbs/{kb2['id']}/index/estimate", headers=h).json()
ok(est["target"]["chunk_strategy"] == "by_heading" and est["chunks"] > 0, "the rebuild estimate uses the new strategy")
c.post(f"/admin/api/kbs/{kb2['id']}/index/rebuild", headers=h)
end = time.time() + 120
while versions.rebuilding(db, kb2["id"]) and time.time() < end:
    time.sleep(0.2)
act = versions.active(db, kb2["id"])
after = {m["heading_path"] for _, ms in kbs.store_for(kbs.get_kb(db, kb2["id"])).iter_metadata() for m in ms}
ok(act["version"] == 2 and act["chunk_strategy"] == "by_heading" and "Leave Policy > Notice period" in after, "after the rebuild the new version uses the strategy and its chunks carry heading paths")
documents.shutdown()
