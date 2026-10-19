"""Read a file into [(page, text)]. Page is the PDF page or slide number; 0 for formats without pages."""
import re
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".pptx", ".md", ".txt"}
OLD_FORMATS = {".doc": ".docx", ".ppt": ".pptx"}


def load_pages(path: Path) -> list[tuple[int, str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [(0, "\n".join(parts))]
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        texts.append(" | ".join(c.text for c in row.cells if c.text.strip()))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append(slide.notes_slide.notes_text_frame.text)
            out.append((i, "\n".join(texts)))
        return out
    if ext in (".md", ".txt"):
        return [(0, path.read_text(encoding="utf-8", errors="ignore"))]
    raise ValueError(f"{ext or 'This'} files are not supported.")


# ---- structure-aware reading (used by the chunking strategies other than fixed size) -------------------------------
from dataclasses import dataclass


@dataclass(frozen=True)
class Block:
    page: int            # PDF page or slide number; 0 for formats without pages
    kind: str            # text | heading | table | code
    text: str
    level: int = 0       # heading level (1 = top)


def _paragraphs(text: str, page: int) -> list[Block]:
    return [Block(page, "text", p.strip()) for p in re.split(r"\n\s*\n", text) if p.strip()]


def _rows(rows) -> str:
    return "\n".join(" | ".join((c or "").replace("\n", " ").strip() for c in r) for r in rows if any((c or "").strip() for c in r))


def load_blocks(path: Path) -> list[Block]:
    """Read a file as blocks: headings, text paragraphs, tables and code blocks, each with its page or slide."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf_blocks(path)
    if ext == ".docx":
        return _docx_blocks(path)
    if ext == ".pptx":
        return _pptx_blocks(path)
    if ext == ".md":
        return _markdown_blocks(path.read_text(encoding="utf-8", errors="ignore"))
    if ext == ".txt":
        return _paragraphs(path.read_text(encoding="utf-8", errors="ignore"), 0)
    raise ValueError(f"{ext or 'This'} files are not supported.")


def _pdf_blocks(path: Path) -> list[Block]:
    try:
        import pdfplumber
        out: list[Block] = []
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                tables = page.find_tables()
                boxes = [t.bbox for t in tables]
                body = page.filter(lambda o: not any(b[0] - 1 <= o.get("x0", 0) and o.get("x1", 0) <= b[2] + 1 and b[1] - 1 <= o.get("top", 0)
                                                     and o.get("bottom", 0) <= b[3] + 1 for b in boxes)) if boxes else page
                out += _paragraphs(body.extract_text() or "", i)
                for t in tables:
                    text = _rows(t.extract())
                    if text:
                        out.append(Block(i, "table", text))
        return out
    except Exception:
        return [b for p, t in load_pages(path) for b in _paragraphs(t, p)]      # fall back to plain page text


def _docx_blocks(path: Path) -> list[Block]:
    from docx import Document
    from docx.table import Table
    out: list[Block] = []
    for item in Document(str(path)).iter_inner_content():
        if isinstance(item, Table):
            text = _rows([[c.text for c in r.cells] for r in item.rows])
            if text:
                out.append(Block(0, "table", text))
            continue
        text = item.text.strip()
        if not text:
            continue
        style = (item.style.name if item.style is not None else "") or ""
        m = re.match(r"Heading\s*(\d)", style, re.I)
        if m or style.lower() == "title":
            out.append(Block(0, "heading", text, int(m.group(1)) if m else 1))
        else:
            out.append(Block(0, "text", text))
    return out


def _pptx_blocks(path: Path) -> list[Block]:
    from pptx import Presentation
    out: list[Block] = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        if title:
            out.append(Block(i, "heading", title, 1))
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                text = _rows([[c.text for c in r.cells] for r in shape.table.rows])
                if text:
                    out.append(Block(i, "table", text))
            elif shape.has_text_frame and shape.text_frame.text.strip() and shape != slide.shapes.title:
                out.append(Block(i, "text", shape.text_frame.text.strip()))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(Block(i, "text", slide.notes_slide.notes_text_frame.text.strip()))
    return out


def _markdown_blocks(text: str) -> list[Block]:
    out: list[Block] = []
    para: list[str] = []
    lines = text.splitlines()
    i = 0

    def flush():
        if para:
            out.append(Block(0, "text", "\n".join(para).strip()))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            flush()
            code = [line]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            code.append(lines[i] if i < len(lines) else "```")
            out.append(Block(0, "code", "\n".join(code)))
        elif re.match(r"^#{1,6}\s+\S", line):
            flush()
            m = re.match(r"^(#{1,6})\s+(.*)$", line)
            out.append(Block(0, "heading", m.group(2).strip(), len(m.group(1))))
        elif line.strip().startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            i -= 1
            out.append(Block(0, "table", _rows(rows)))
        elif not line.strip():
            flush()
        else:
            para.append(line)
        i += 1
    flush()
    return [b for b in out if b.text.strip()]
